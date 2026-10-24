"""Clubs trouvés par nous-mêmes (répertoire SIRENE) et vérifiés un par un sur leur page officielle.

Méthode : le répertoire SIRENE (Licence Ouverte) sert seulement à découvrir le nom d'un club. Chaque fiche est
ensuite remplie à partir de ce que le club publie lui-même (son site, sa page HelloAsso ou JeVeuxAider), indiquée
dans « verifie_sur ». Un club n'est importé que s'il dit lui-même accueillir des personnes en situation de handicap.
L'adresse du siège donnée par SIRENE n'est pas reprise (c'est souvent le domicile d'un bénévole).

Le fichier Excel de suivi (statut de chaque club examiné, raison d'un refus) est produit avec --excel.

Utilisation : python clubs_verifies.py > 13_import_clubs_verifies.sql
              python clubs_verifies.py --excel suivi.xlsx
"""
import sys

from csv_vers_sql import requete

SOURCE = "Vérifié sur la page officielle du club (découvert via le répertoire SIRENE)"
VERIFIE_LE = "2026-10-24"

# Clubs retenus : chaque information vient de la page indiquée dans « verifie_sur »
CLUBS = [
    {"siren": "404376873", "nom": "SASEP – Sport Adapté du Sud-Est de Paris", "sport": "natation",
     "adresse": "Piscine de la Butte-aux-Cailles", "code_postal": "75013", "ville": "Paris 13e",
     "handicaps": "mental", "pratique": "sport adapte",
     "accessibilite": "Jeunes de 6 à 21 ans, encadrés par des psychomotriciens et des bénévoles",
     "telephone": "07 83 90 56 62", "email": "associationsasep@gmail.com",
     "site": "https://www.jeveuxaider.gouv.fr/organisations/18654-sasep-sport-adapte-du-sud-est-de-paris",
     "verifie_sur": "https://www.jeveuxaider.gouv.fr/organisations/18654-sasep-sport-adapte-du-sud-est-de-paris"},
    {"siren": "899929541", "nom": "Paris Basket Fauteuil", "sport": "basket fauteuil",
     "code_postal": "75018", "ville": "Paris 18e", "handicaps": "moteur", "pratique": "handisport",
     "accessibilite": "Tous niveaux, des débutants aux compétiteurs", "deja": True,
     "site": "https://www.helloasso.com/associations/paris-basket-fauteuil",
     "verifie_sur": "https://www.helloasso.com/associations/paris-basket-fauteuil"},
    {"siren": "883070518", "nom": "Club Sportif des Sourds-Muets de Paris – Football", "sport": "football",
     "code_postal": "75019", "ville": "Paris 19e", "handicaps": "auditif", "pratique": "handisport",
     "site": "https://www.helloasso.com/associations/club-sportif-des-sourds-muets-de-paris-section-football",
     "verifie_sur": "https://www.helloasso.com/associations/club-sportif-des-sourds-muets-de-paris-section-football"},
]

# Clubs examinés mais non importés, avec la raison
ECARTES = [
    ("441806791", "Aéro-Club des Sourds de France", "Paris 16e", "À confirmer",
     "Existe (page HelloAsso, vols ULM), mais le lieu de pratique n'est pas confirmé par le club : Coulommiers selon un article de 2016"),
    ("450124938", "Club Sportif des Sourds Muets de Paris (CSSMP)", "Paris 19e", "À confirmer",
     "Association multisport : aucune page officielle trouvée qui décrive ses activités actuelles (seule la section football en a une)"),
    ("428429013", "Tennis Club des Sourds de Paris", "Paris 15e", "Écarté", "Aucune page officielle trouvée"),
    ("434581435", "Handisport Organisation Promotion (HOP)", "Paris 20e", "Écarté",
     "Organise une compétition (Handisport Open Paris), ce n'est pas un club"),
    ("924801202", "Rugby Fauteuil France Développement Compétition", "Paris 15e", "Écarté",
     "Association de soutien (matériel, compétitions), ce n'est pas un club"),
    ("813274867", "Handivoile", "Paris 8e", "Écarté", "Organise une journée nationale d'initiation, ce n'est pas un club"),
    ("928636869", "Handisport pour tous", "Paris 12e", "Écarté", "Aucune page officielle trouvée"),
    ("891532202", "Sport Santé Adapt", "Paris 20e", "Écarté", "Aucune page officielle trouvée"),
    ("784670713", "Ass Sport Handi Moteur Ile Fr", "Paris 12e", "Écarté", "Aucune page officielle trouvée"),
    ("992149500", "AP2S – Pratique sécurisée du parasport et du sport", "Paris 11e", "Écarté",
     "Créée en septembre 2025, aucune page officielle trouvée"),
    ("520255720", "Relais Sport Handicap", "Paris 20e", "Écarté", "Aucune page officielle trouvée"),
]


def fiches():
    for c in CLUBS:
        if c.get("deja"):                                            # déjà sur Paralliance : rien à importer
            continue
        yield {k: v for k, v in c.items() if k not in ("siren", "verifie_sur", "deja")} | {"verifie_le": VERIFIE_LE, "source": SOURCE}


def excel(chemin):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    from openpyxl.utils import get_column_letter
    wb = Workbook()
    ws = wb.active
    ws.title = "Clubs examinés"
    titres = ["Statut", "SIREN", "Nom", "Ville", "Sports", "Handicaps", "Pratique", "Téléphone", "Mail",
              "Page officielle vérifiée", "Raison / remarque", "Vérifié le"]
    ws.append(titres)
    for c in CLUBS:
        ws.append(["Déjà sur le site" if c.get("deja") else "Importé", c["siren"], c["nom"], c["ville"], c["sport"], c["handicaps"], c["pratique"],
                   c.get("telephone", ""), c.get("email", ""), c["verifie_sur"], c.get("accessibilite", ""), VERIFIE_LE])
    for siren, nom, ville, statut, raison in ECARTES:
        ws.append([statut, siren, nom, ville, "", "", "", "", "", "", raison, VERIFIE_LE])
    couleurs = {"Importé": "D9EAD3", "Déjà sur le site": "CFE2F3", "À confirmer": "FFF2CC", "Écarté": "F4CCCC"}
    for ligne in ws.iter_rows(min_row=2):
        ligne[0].fill = PatternFill("solid", fgColor=couleurs[ligne[0].value])
    for cellule in ws[1]:
        cellule.font = Font(bold=True)
    largeurs = [12, 11, 45, 14, 18, 12, 14, 15, 30, 60, 70, 11]
    for i, l in enumerate(largeurs, 1):
        ws.column_dimensions[get_column_letter(i)].width = l
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    wb.save(chemin)


if __name__ == "__main__":
    if len(sys.argv) > 2 and sys.argv[1] == "--excel":
        excel(sys.argv[2])
    else:
        print(requete(list(fiches()), SOURCE), end="")
